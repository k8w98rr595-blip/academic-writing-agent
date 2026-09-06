"""Synthetic, credential-free regressions for saved versions and batch previews."""
import io

import pytest
from docx import Document as WordDocument
from docx.oxml import OxmlElement
from fastapi import HTTPException
from sqlalchemy import select

from services.api.app.database import session_scope
from services.api.app.documents import DOCX_MIME, extract_docx_text
from services.api.app.models import Document, PatchRecord, RewriteSession


def word_bytes(word):
    stream = io.BytesIO()
    word.save(stream)
    return stream.getvalue()


def test_word_preserves_interleaved_table_order():
    word = WordDocument()
    word.add_paragraph("Before table")
    word.add_table(rows=1, cols=2).cell(0, 0).text = "Middle cell"
    word.add_paragraph("After table")
    text = extract_docx_text(word_bytes(word))
    assert text.index("Before table") < text.index("Middle cell") < text.index("After table")


def test_word_numbering_inherited_from_style_is_not_silently_lost():
    word = WordDocument()
    word.add_paragraph("Automatically numbered evidence", style="List Number")
    with pytest.raises(HTTPException) as failure:
        extract_docx_text(word_bytes(word))
    assert "自动编号" in failure.value.detail


@pytest.mark.parametrize("tag", ["m:oMath", "w:footnoteReference", "w:endnoteReference", "w:drawing", "w:ins", "w:del", "w:sdt", "w:fldChar", "w:numPr"])
def test_word_unsupported_content_fails_closed(tag):
    word = WordDocument()
    word.add_paragraph("Synthetic test prose")._p.append(OxmlElement(tag))
    with pytest.raises(HTTPException) as failure:
        extract_docx_text(word_bytes(word))
    assert failure.value.status_code == 422
    assert "不支持" in failure.value.detail


@pytest.mark.parametrize("kind", ["merged", "nested", "header"])
def test_word_rejects_lossy_structures(kind):
    word = WordDocument()
    word.add_paragraph("Synthetic test prose")
    if kind == "header":
        word.sections[0].header.paragraphs[0].text = "Important header content"
    else:
        table = word.add_table(rows=2, cols=2)
        if kind == "merged":
            table.cell(0, 0).merge(table.cell(0, 1))
        else:
            table.cell(0, 0).add_table(rows=1, cols=1)
    with pytest.raises(HTTPException) as failure:
        extract_docx_text(word_bytes(word))
    assert failure.value.status_code == 422


@pytest.fixture
def preview(client, headers):
    sentence = "It is important to note that careful evidence supports responsible analysis and transparent reasoning in this synthetic exercise. "
    text = "\n\n".join(sentence * 12 for _ in range(3))
    created = client.post("/api/v1/documents", headers=headers, data={"title": "Synthetic reliability test", "text": text})
    assert created.status_code == 201
    document = created.json()["document"]
    assert client.post(f"/api/v1/documents/{document['id']}/analyses", headers=headers).status_code == 201
    response = client.post(f"/api/v1/documents/{document['id']}/first-pass-rewrite", headers=headers,
                           json={"version_id": document["currentVersion"]["id"]})
    assert response.status_code == 201
    result = response.json()
    assert len(result["document"]["patches"]) == 3
    return result


def decide(client, headers, preview, ids):
    return client.post(f"/api/v1/rewrite-sessions/{preview['rewriteSessionId']}/batch-decision", headers=headers,
        json={"expected_base_version_id": preview["document"]["currentVersion"]["id"], "accepted_patch_ids": ids})


def test_batch_partial_accept_restore_export_and_delete(client, headers, preview):
    document = preview["document"]
    base = document["currentVersion"]
    patch = document["patches"][0]
    response = decide(client, headers, preview, [patch["id"]])
    assert response.status_code == 200
    revised = response.json()["document"]
    assert revised["currentVersion"]["number"] == 2
    assert revised["analysis"]["isStale"] is True
    assert sum(p["status"] == "accepted" for p in revised["patches"]) == 1
    assert sum(p["status"] == "rejected" for p in revised["patches"]) == 2
    for original, current in zip(base["paragraphs"], revised["currentVersion"]["paragraphs"]):
        assert current["text"] == (patch["revisedText"] if current["id"] == patch["paragraphId"] else original["text"])
    # Decisions do not generate another paid proposal or detection.
    assert decide(client, headers, preview, [patch["id"]]).status_code == 409
    exported = client.post(f"/api/v1/documents/{document['id']}/exports", headers=headers,
                          json={"expected_version_id": revised["currentVersion"]["id"]})
    assert exported.status_code == 200
    assert patch["revisedText"] in extract_docx_text(exported.content)
    restored = client.post(f"/api/v1/documents/{document['id']}/versions/{base['id']}/restore", headers=headers,
                           json={"expected_current_version_id": revised["currentVersion"]["id"]})
    assert restored.status_code == 200
    assert restored.json()["document"]["currentVersion"]["paragraphs"] == base["paragraphs"]
    assert client.delete(f"/api/v1/documents/{document['id']}", headers=headers).status_code == 204
    with session_scope() as db:
        assert db.scalar(select(RewriteSession).where(RewriteSession.document_id == document["id"])) is None
        assert db.scalar(select(PatchRecord).where(PatchRecord.document_id == document["id"])) is None


def test_batch_reject_leaves_version_and_analysis_unchanged(client, headers, preview):
    response = decide(client, headers, preview, [])
    assert response.status_code == 200
    document = response.json()["document"]
    assert document["currentVersion"]["number"] == 1
    assert document["analysis"]["isStale"] is False
    assert all(p["status"] == "rejected" for p in document["patches"])


def test_batch_preview_survives_reload_with_truthful_provider_label(client, headers, preview):
    doc = client.get(f"/api/v1/documents/{preview['document']['id']}", headers=headers).json()["document"]
    assert len(doc["patches"]) == 3
    assert all(p["isMock"] and p["batch"] and p["status"] == "pending" for p in doc["patches"])


@pytest.mark.parametrize("choice", ["foreign", "duplicate", "stale", "protected", "owner"])
def test_batch_invalid_decisions_cannot_change_document(client, headers, preview, choice):
    doc = preview["document"]
    patch = doc["patches"][0]
    ids = [patch["id"]]
    expected = 422
    if choice == "foreign":
        ids = ["patch_other_document"]
    elif choice == "duplicate":
        ids *= 2
    elif choice == "stale":
        saved = client.patch(f"/api/v1/documents/{doc['id']}", headers=headers,
            json={"base_version_id": doc["currentVersion"]["id"], "paragraphs": doc["currentVersion"]["paragraphs"]})
        assert saved.status_code == 200
        expected = 409
    elif choice == "protected":
        with session_scope() as db:
            db.get(PatchRecord, patch["id"]).revised_text += " A fabricated 42% claim."
        expected = 409
    else:
        with session_scope() as db:
            db.get(Document, doc["id"]).owner_email = "other@example.invalid"
        expected = 404
    assert decide(client, headers, preview, ids).status_code == expected
    with session_scope() as db:
        assert db.get(PatchRecord, patch["id"]).status == "pending"


def test_export_requires_expected_current_version(client, headers, preview):
    document = preview["document"]
    base = document["currentVersion"]
    edited = [dict(p) for p in base["paragraphs"]]
    edited[0]["text"] += " Newly saved text."
    saved = client.patch(f"/api/v1/documents/{document['id']}", headers=headers,
                        json={"base_version_id": base["id"], "paragraphs": edited}).json()["document"]
    assert client.post(f"/api/v1/documents/{document['id']}/exports", headers=headers,
                       json={"expected_version_id": base["id"]}).status_code == 409
    exported = client.post(f"/api/v1/documents/{document['id']}/exports", headers=headers,
                           json={"expected_version_id": saved["currentVersion"]["id"]})
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(DOCX_MIME)
    assert "Newly saved text." in extract_docx_text(exported.content)


def test_batch_route_requires_authentication(client):
    assert client.post("/api/v1/rewrite-sessions/rewrite_unknown/batch-decision",
                       json={"expected_base_version_id": "version_unknown", "accepted_patch_ids": []}).status_code == 401


def test_large_pending_batch_is_not_truncated_to_history_limit(client, headers):
    sentence = "It is important to note that this synthetic argument links evidence with responsible interpretation and clear reasoning."
    document = client.post("/api/v1/documents", headers=headers,
        data={"title": "Synthetic long preview", "text": "\n\n".join([sentence] * 40)}).json()["document"]
    assert client.post(f"/api/v1/documents/{document['id']}/analyses", headers=headers).status_code == 201
    preview = client.post(f"/api/v1/documents/{document['id']}/first-pass-rewrite", headers=headers,
        json={"version_id": document["currentVersion"]["id"]}).json()
    assert len(preview["document"]["patches"]) == 40
    refreshed = client.get(f"/api/v1/documents/{document['id']}", headers=headers).json()["document"]
    assert len([p for p in refreshed["patches"] if p["batch"] and p["status"] == "pending"]) == 40
