"""Credential-free launch regressions. All users and writing are synthetic."""
from dataclasses import replace
import io

import pytest
from docx import Document as WordDocument
from fastapi import HTTPException
from sqlalchemy import func, select

from services.api.app import accounts, billing, billing_catalog, main, security, provider_usage
from services.api.app.config import get_settings
from services.api.app.database import session_scope
from services.api.app.documents import build_docx
from services.api.app.models import AnalysisRun, BillingAccount, Document, DocumentVersion, PatchRecord, ProductUsageReservation, ProviderUsageEvent, RewriteSession, UserAccount
from services.api.app.text import assert_protected_equal, sentence_ranges, text_metrics, validate_english_coursework
from tests.test_workflow import create_document

PASSWORD = "synthetic password for tests"


@pytest.fixture
def public_settings(monkeypatch):
    configured = replace(get_settings(), public_registration_enabled=True, public_ai_enabled=False,
                         public_launch_acknowledged=True, operator_name="Synthetic operator", support_email="test@example.com")
    for module in (accounts, billing, billing_catalog, security, provider_usage):
        monkeypatch.setattr(module, "get_settings", lambda: configured)
    monkeypatch.setattr(main, "settings", configured)
    return configured


def register(client, username="student_one", **changes):
    return client.post("/api/v1/auth/register", json={"username": username, "password": PASSWORD,
        "accepted_terms_version": accounts.TERMS_VERSION, "adult_confirmed": True, "cloud_storage_consent": True, **changes})


def auth(payload):
    return {"Authorization": f"Bearer {payload['session_token']}"}


def test_public_default_is_local_only(client):
    result = client.get("/api/v1/public/config")
    assert result.status_code == 200
    assert result.json()["localExperienceAvailable"] is True
    assert result.json()["registrationEnabled"] is False
    assert result.json()["publicAiEnabled"] is False
    assert register(client).status_code == 403
    assert client.get("/api/v1/account").status_code == 401


@pytest.mark.parametrize("changes", [{"adult_confirmed": False}, {"cloud_storage_consent": False}, {"accepted_terms_version": "old"}, {"role": "owner"}, {"provider_processing_consent": True}])
def test_registration_rejects_missing_consent_and_client_privilege(client, public_settings, changes):
    assert register(client, **changes).status_code == 422
    with session_scope() as db:
        assert db.scalar(select(func.count(UserAccount.id))) == 0


def test_account_hashes_and_duplicate_username(client, public_settings):
    response = register(client)
    assert response.status_code == 201
    payload = response.json()
    assert client.get("/api/v1/account", headers=auth(payload)).json()["role"] == "student"
    assert register(client, "STUDENT_ONE").status_code == 409
    assert register(client, "admin").status_code == 409
    with session_scope() as db:
        user = db.scalar(select(UserAccount))
        assert user.password_hash != PASSWORD and user.password_hash.startswith("$argon2id$")
        assert user.recovery_hash != payload["recovery_key"]
        assert user.provider_consent is False
    login = client.post("/api/v1/auth/login", json={"email": "STUDENT_ONE", "password": PASSWORD})
    assert login.status_code == 200
    assert client.get("/api/v1/provider-usage/summary", headers=auth(login.json())).status_code == 403


def test_registration_capacity_and_rate_limits(client, public_settings, monkeypatch):
    configured = replace(public_settings, public_max_accounts=1)
    monkeypatch.setattr(accounts, "get_settings", lambda: configured)
    assert register(client).status_code == 201
    assert register(client, "student_two").status_code == 503
    for _ in range(3):
        assert register(client, "student_three").status_code == 503
    assert register(client, "student_four").status_code == 429


def test_student_document_capacity_enforced_with_billing_off(client, public_settings, coursework_text):
    headers = auth(register(client).json())
    for _ in range(public_settings.billing_free_documents):
        create_document(client, headers, coursework_text)
    result = client.post("/api/v1/documents", headers=headers, data={"title": "Beyond capacity", "text": coursework_text})
    assert result.status_code == 402
    assert client.get("/api/v1/billing/summary", headers=headers).json()["plan"]["key"] == "free"


def test_cross_account_document_patch_job_and_export_isolation(client, public_settings, coursework_text):
    first, second = auth(register(client).json()), auth(register(client, "student_two").json())
    document = create_document(client, first, coursework_text)
    document_id = document["id"]
    assert client.get("/api/v1/documents", headers=second).json()["documents"] == []
    assert client.get(f"/api/v1/documents/{document_id}", headers=second).status_code == 404
    assert client.delete(f"/api/v1/documents/{document_id}", headers=second).status_code == 404
    assert client.post(f"/api/v1/documents/{document_id}/exports", headers=second).status_code == 404
    analyzed = client.post(f"/api/v1/documents/{document_id}/analyses", headers=first).json()
    assert client.get(f"/api/v1/jobs/{analyzed['jobId']}/events", headers=second).status_code == 404
    session = client.post(f"/api/v1/documents/{document_id}/rewrite-sessions", headers=first, json={"version_id": document["currentVersion"]["id"]}).json()["rewriteSession"]
    paragraph = document["currentVersion"]["paragraphs"][1]
    message = {"paragraph_id": paragraph["id"], "instruction": "Improve clarity"}
    assert client.post(f"/api/v1/rewrite-sessions/{session['id']}/messages", headers=second, json=message).status_code == 404
    patch = client.post(f"/api/v1/rewrite-sessions/{session['id']}/messages", headers=first, json=message).json()["patch"]
    assert client.post(f"/api/v1/patches/{patch['id']}/accept", headers=second, json={"expected_base_version_id": patch["baseVersionId"]}).status_code == 404
    assert client.patch(f"/api/v1/documents/{document_id}", headers=second, json={"base_version_id": document["currentVersion"]["id"], "paragraphs": document["currentVersion"]["paragraphs"]}).status_code == 404


def test_real_student_ai_closed_before_any_provider_or_reservation(client, public_settings, coursework_text, monkeypatch):
    headers = auth(register(client).json())
    document = create_document(client, headers, coursework_text)
    monkeypatch.setattr(main, "settings", replace(public_settings, detector_mode="pangram", rewrite_mode="deepseek"))
    def forbidden(*args, **kwargs):
        pytest.fail("Provider must not be called while public AI is closed")
    monkeypatch.setattr(main, "run_detection", forbidden)
    monkeypatch.setattr(main, "propose_rewrite", forbidden)
    assert client.post(f"/api/v1/documents/{document['id']}/analyses", headers=headers).status_code == 403
    assert client.post(f"/api/v1/documents/{document['id']}/first-pass-rewrite", headers=headers, json={"version_id": document["currentVersion"]["id"]}).status_code == 403
    session = client.post(f"/api/v1/documents/{document['id']}/rewrite-sessions", headers=headers, json={"version_id": document["currentVersion"]["id"]}).json()["rewriteSession"]
    assert client.post(f"/api/v1/rewrite-sessions/{session['id']}/messages", headers=headers, json={"paragraph_id": document["currentVersion"]["paragraphs"][1]["id"], "instruction": "Improve clarity"}).status_code == 403
    with session_scope() as db:
        assert db.scalar(select(func.count(ProviderUsageEvent.id))) == 0


def test_provider_consent_separate_versioned_and_revocable(client, public_settings, monkeypatch):
    response = register(client).json()
    headers = auth(response)
    consent = {"accepted_terms_version": accounts.TERMS_VERSION, "consent": True}
    assert client.post("/api/v1/account/provider-consent", headers=headers, json=consent).status_code == 403
    configured = replace(public_settings, public_ai_enabled=True)
    monkeypatch.setattr(accounts, "get_settings", lambda: configured)
    with session_scope() as db:
        principal = db.scalar(select(UserAccount.principal))
    with pytest.raises(HTTPException):
        accounts.require_public_ai(principal, "pangram")
    assert client.post("/api/v1/account/provider-consent", headers=headers, json={**consent, "accepted_terms_version": "old"}).status_code == 403
    assert client.post("/api/v1/account/provider-consent", headers=headers, json=consent).status_code == 200
    accounts.require_public_ai(principal, "pangram")
    assert client.post("/api/v1/account/provider-consent", headers=headers, json={**consent, "consent": False}).status_code == 200
    with pytest.raises(HTTPException):
        accounts.require_public_ai(principal, "deepseek")


def test_recovery_rotates_code_and_revokes_old_sessions(client, public_settings):
    payload = register(client).json()
    recovered = client.post("/api/v1/auth/recover", json={"username": "student_one", "recovery_key": payload["recovery_key"], "new_password": PASSWORD + " new"})
    assert recovered.status_code == 200
    assert recovered.json()["recovery_key"] != payload["recovery_key"]
    assert client.get("/api/v1/documents", headers=auth(payload)).status_code == 401
    assert client.post("/api/v1/auth/recover", json={"username": "student_one", "recovery_key": payload["recovery_key"], "new_password": PASSWORD}).status_code == 401
    assert client.post("/api/v1/auth/login", json={"email": "student_one", "password": PASSWORD + " new"}).status_code == 200


def test_password_change_and_account_erasure(client, public_settings, coursework_text):
    payload = register(client).json()
    headers = auth(payload)
    document = create_document(client, headers, coursework_text)
    assert client.post(f"/api/v1/documents/{document['id']}/analyses", headers=headers).status_code == 201
    changed = client.post("/api/v1/account/password", headers=headers, json={"current_password": PASSWORD, "new_password": PASSWORD + " new"})
    assert changed.status_code == 200
    assert client.get("/api/v1/account", headers=headers).status_code == 401
    new_login = client.post("/api/v1/auth/login", json={"email": "student_one", "password": PASSWORD + " new"}).json()
    headers = auth(new_login)
    assert client.request("DELETE", "/api/v1/account", headers=headers, json={"password": PASSWORD}).status_code == 403
    assert client.request("DELETE", "/api/v1/account", headers=headers, json={"password": PASSWORD + " new"}).status_code == 204
    assert client.get("/api/v1/account", headers=headers).status_code == 401
    with session_scope() as db:
        for model in (UserAccount, Document, DocumentVersion, AnalysisRun, PatchRecord, RewriteSession):
            assert db.scalar(select(func.count()).select_from(model)) == 0
        assert db.scalar(select(func.count(BillingAccount.id)).where(BillingAccount.owner_email != public_settings.owner_email)) == 0
        assert db.scalar(select(func.count(ProductUsageReservation.id))) == 0


def test_owner_credentials_cannot_be_mutated_through_student_api(client, headers):
    assert client.post("/api/v1/account/password", headers=headers, json={"current_password": "correct horse battery staple", "new_password": PASSWORD}).status_code == 403
    assert client.request("DELETE", "/api/v1/account", headers=headers, json={"password": "correct horse battery staple"}).status_code == 403
    assert client.get("/api/v1/account", headers=headers).json()["role"] == "owner"


def test_validation_does_not_echo_password_or_recovery_input(client):
    secret = "never-echo-synthetic-input"
    result = client.post("/api/v1/auth/register", json={"password": secret, "recovery_key": secret})
    assert result.status_code == 422 and secret not in result.text


def test_chinese_boundary_and_protected_data():
    assert validate_english_coursework("中" * 500) == 500
    assert text_metrics("中" * 500 + "AI 2026 don't") == {"language": "zh", "count": 503, "unit": "字/词"}
    for text in ("中" * 499, "中" * 12001):
        with pytest.raises(HTTPException):
            validate_english_coursework(text)
    assert len(sentence_ranges("论点。证据！结论？")) == 3
    assert_protected_equal("根据《课程说明》2026年30份样本。", "2026年30份样本来自《课程说明》。")
    with pytest.raises(HTTPException):
        assert_protected_equal("2026年30份样本。", "2026年300份样本。")


def test_chinese_mock_closure_export_and_delete(client, headers):
    paragraph = "值得注意的是，课程论文的论点需要具体证据支持，作者应核验原始材料，不能用工具建议替代独立判断。"
    document = create_document(client, headers, "课程写作\n\n" + paragraph * 15)
    assert document["currentVersion"]["language"] == "zh"
    before = client.post(f"/api/v1/documents/{document['id']}/analyses", headers=headers).json()["analysis"]
    assert before["result"]["isMock"] is True
    session = client.post(f"/api/v1/documents/{document['id']}/rewrite-sessions", headers=headers, json={"version_id": document["currentVersion"]["id"]}).json()["rewriteSession"]
    patch_result = client.post(f"/api/v1/rewrite-sessions/{session['id']}/messages", headers=headers, json={"paragraph_id": document["currentVersion"]["paragraphs"][1]["id"], "instruction": "保持原意，改善表达清晰度"})
    assert patch_result.status_code == 201
    patch = patch_result.json()["patch"]
    accepted = client.post(f"/api/v1/patches/{patch['id']}/accept", headers=headers, json={"expected_base_version_id": patch["baseVersionId"]})
    assert accepted.status_code == 200 and accepted.json()["document"]["analysis"]["isStale"] is True
    assert client.post(f"/api/v1/documents/{document['id']}/analyses", headers=headers).status_code == 201
    export = client.post(f"/api/v1/documents/{document['id']}/exports", headers=headers)
    assert export.status_code == 200 and "需要注意" in WordDocument(io.BytesIO(export.content)).paragraphs[1].text
    assert client.delete(f"/api/v1/documents/{document['id']}", headers=headers).status_code == 204


def test_docx_ai_disclosure_does_not_change_body():
    data = build_docx("Synthetic", [{"id": "p1", "text": "合成正文。"}], 2, ai_assisted=True)
    exported = WordDocument(io.BytesIO(data))
    assert exported.paragraphs[0].text == "合成正文。"
    assert "AI 辅助写作" in exported.sections[0].footer.paragraphs[0].text
    assert "AI 辅助写作" in exported.core_properties.comments


def test_global_paid_budget_cannot_be_bypassed_by_second_user(monkeypatch):
    configured = replace(get_settings(), global_paid_hourly_limit=1)
    monkeypatch.setattr(provider_usage, "get_settings", lambda: configured)
    spec = [provider_usage.ProviderCallSpec("rewrite", "synthetic-model", "synthetic-operation")]
    provider_usage.reserve_provider_calls("student_one", "DeepSeek", spec)
    with pytest.raises(HTTPException) as raised:
        provider_usage.reserve_provider_calls("student_two", "DeepSeek", spec)
    assert raised.value.status_code == 429


@pytest.mark.parametrize("variable,value", [("PUBLIC_REGISTRATION_ENABLED", "1"), ("PUBLIC_AI_ENABLED", "1"), ("PUBLIC_MAX_ACCOUNTS", "0"), ("GLOBAL_PAID_HOURLY_LIMIT", "99999")])
def test_launch_settings_fail_closed(monkeypatch, variable, value):
    monkeypatch.setenv(variable, value)
    get_settings.cache_clear()
    try:
        with pytest.raises(RuntimeError):
            get_settings()
    finally:
        get_settings.cache_clear()
